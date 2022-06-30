from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

wb = load_workbook("./rpabasic/excel/range.xlsx")
ws = wb.active

# 막대 차트
# 차트 범위 설정
bar_value = Reference(ws, min_row=2, max_row=11, min_col=3, max_col=4)

# 차트 종류 설정
bar_chart = BarChart()
bar_chart.add_data(bar_value)

# 차트 삽입
ws.add_chart(bar_chart, "E1")

wb.save("./rpabasic/excel/range.xlsx")
